import base64
from io import BytesIO
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font
from datetime import date, datetime
from odoo import api, fields, models
from odoo.exceptions import ValidationError
from odoo.tools import html2plaintext
from odoo.addons.account_withholding_tax_on_payment.tools.date_tools import (
    get_locale_date_format,
)
from odoo.addons.account_withholding_tax_on_payment.tools.get_module_resource import (
    get_module_resource,
)


class AccountWithholdingTax(models.Model):
    _name = "account.withholding.tax"
    _description = "Account Withholding Tax"

    @api.model
    def _get_wht_payment(self):
        return [
            ("wht", "(1) หัก ณ ที่จ่าย"),
            ("forever", "(2) ออกให้ตลอดไป"),
            ("once", "(3) ออกครั้งเดียว"),
            ("other", "(4) อื่น ๆ"),
        ]

    name = fields.Char(
        copy=False,
        readonly=True,
        default="/",
    )
    wht_payment = fields.Selection(
        selection=_get_wht_payment,
        default="wht",
        string="WHT Payment",
        required=True,
    )
    partner_id = fields.Many2one(
        comodel_name="res.partner",
        readonly=True,
    )
    payment_id = fields.Many2one(
        comodel_name="account.payment",
        ondelete="restrict",
        copy=False,
        index=True,
    )
    company_id = fields.Many2one(
        comodel_name="res.company",
        required=True,
        readonly=True,
        default=lambda self: self.env.company,
    )
    currency_id = fields.Many2one(
        related="company_id.currency_id",
    )
    document_date = fields.Date(
        required=True,
        readonly=True,
    )
    state = fields.Selection(
        selection=[
            ("draft", "Draft"),
            ("done", "Done"),
        ],
        default="draft",
        copy=False,
    )
    source_document = fields.Char(
        copy=False,
    )
    remark = fields.Char(
        copy=False,
    )
    withholding_line_ids = fields.One2many(
        comodel_name="account.withholding.tax.line",
        inverse_name="withholding_tax_id",
    )
    total_withholding_amount = fields.Monetary(
        compute="_compute_total_withholding_amount",
        string="Total Withholding Amount",
        currency_field="currency_id",
    )
    total_base_amount = fields.Monetary(
        compute="_compute_total_base_amount",
        string="Total Base Amount",
        currency_field="currency_id",
    )
    pnd_type = fields.Selection(
        selection=[
            ("pnd3", "P.N.D.3"),
            ("pnd53", "P.N.D.53"),
        ],
        compute="_compute_pnd_type",
        string="Filing According",
        store=True,
    )
    pnd_sequence = fields.Char(
        copy=False,
        readonly=True,
    )
    move_ids = fields.Many2many(
        comodel_name="account.move",
        relation="account_withholding_tax_account_move_rel",
        column1="withholding_tax_id",
        column2="move_id",
        string="Journal Entries",
        copy=False,
    )
    is_stamp_signature = fields.Boolean(
        string="Stamp Signature",
        default=False,
    )

    @api.depends("withholding_line_ids.tax_id.invoice_repartition_line_ids.account_id")
    def _compute_pnd_type(self):
        pnd_mapping = {
            "+PND3": "pnd3",
            "+PND53": "pnd53",
            "-PND3": "pnd3",
            "-PND53": "pnd53",
        }
        for rec in self:
            tags = rec.withholding_line_ids.mapped(
                "tax_id.invoice_repartition_line_ids.tag_ids.name"
            )
            rec.pnd_type = next(
                (pnd_mapping[t] for t in tags if t in pnd_mapping),
                False,
            )

    @api.depends("withholding_line_ids.amount")
    def _compute_total_withholding_amount(self):
        for rec in self:
            rec.total_withholding_amount = round(
                sum(rec.withholding_line_ids.mapped("amount")), 2
            )

    @api.depends("withholding_line_ids.base_amount")
    def _compute_total_base_amount(self):
        for rec in self:
            rec.total_base_amount = round(
                sum(rec.withholding_line_ids.mapped("base_amount")), 2
            )

    @api.onchange("withholding_line_ids")
    def _onchange_withholding_line_ids(self):
        self.ensure_one()
        if not self.withholding_line_ids._need_update_withholding_lines_placeholder():
            return

        self.withholding_line_ids._update_placeholders()

    def _get_wht_data(self):
        vals = []
        for seq in self.withholding_line_ids.mapped("tax_section"):
            seq_lines = self.withholding_line_ids.filtered(
                lambda x, y=seq: x.tax_section == y
            )
            vals.append(
                {
                    "sequence": seq,
                    "base_amount": sum(seq_lines.mapped("base_amount")),
                    "amount": sum(seq_lines.mapped("amount")),
                    "other": ", ".join(
                        seq_lines.filtered(
                            lambda x: x.tax_section in ("425", "6")
                        ).mapped("other")
                        or []
                    ),
                }
            )
        return vals

    def action_set_to_draft(self):
        for rec in self:
            rec.state = "draft"

    def action_done(self):
        for rec in self:
            if rec.state != "draft":
                raise ValidationError(
                    self.env._(
                        "Only Draft document can be confirm, "
                        "Please check your document status and try again."
                    )
                )
            if rec.payment_id.state != "paid":
                raise ValidationError(
                    self.env._(
                        "Payment must be paid, Please check "
                        "your payment and try again."
                    )
                )
            name = rec.name
            if not name or name == "/":
                wht_sequence = rec.company_id.wht_sequence_id
                if not wht_sequence:
                    rec.company_id._create_wht_per_company_sequences()
                    wht_sequence = rec.company_id.wht_sequence_id
                name = wht_sequence.next_by_id(sequence_date=rec.document_date)
            rec.write(
                {
                    "name": name,
                    "state": "done",
                    "pnd_sequence": rec.pnd_sequence
                    or rec._get_pnd_sequence(rec.pnd_type),
                }
            )

    def action_export_xlsx(self):
        self.ensure_one()
        template = get_module_resource(
            "account_withholding_tax_on_payment",
            "static/template",
            "withholding_tax_template.xlsx",
        )
        wb = openpyxl.load_workbook(template)
        ws = wb.active

        self.write_cell(
            ws=ws,
            cell="Q3",
            value=self.name,
        )
        # ผู้ถูกเสียภาษี
        self.write_cell(
            ws=ws,
            cell="C6",
            value=self.company_id.partner_id.display_name,
        )
        self.write_cell(
            ws=ws,
            cell="P5",
            value=self.company_id.vat,
        )
        self.write_cell(
            ws=ws,
            cell="C8",
            value=self.company_id.partner_id._get_address_partner(),
        )

        # ผู้ถูกหักภาษี
        self.write_cell(
            ws=ws,
            cell="C14",
            value=self.partner_id.display_name,
        )
        self.write_cell(
            ws=ws,
            cell="C12",
            value=self.partner_id._get_address_partner(),
        )
        self.write_cell(
            ws=ws,
            cell="P11",
            value=self.partner_id.vat,
        )
        self.write_cell(
            ws=ws,
            cell="D16",
            value=self.pnd_sequence,
        )

        # ภ.ง.ด
        self.write_cell(ws=ws, cell=self._get_pnd_type(self.pnd_type), value="✓")

        for seq in self.withholding_line_ids.mapped("tax_section"):
            seq_lines = self.withholding_line_ids.filtered(
                lambda x, y=seq: x.tax_section == y
            )
            other = ", ".join(
                seq_lines.filtered(lambda x: x.tax_section in ("425", "6")).mapped(
                    "other"
                )
                or []
            )
            self.write_cell(
                ws=ws,
                cell=self._get_tax_section_line(seq).get("date"),
                value=self.get_report_locale_date(
                    dt=self.document_date,
                    py_format=None,
                    be_year=True,
                ),
            )
            self.write_cell(
                ws=ws,
                cell=self._get_tax_section_line(seq).get("base_amount"),
                value=sum(seq_lines.mapped("base_amount")),
            )
            self.write_cell(
                ws=ws,
                cell=self._get_tax_section_line(seq).get("amount"),
                value=sum(seq_lines.mapped("amount")),
            )
            if other:
                self.write_cell(
                    ws=ws,
                    cell=self._get_tax_section_line(seq).get("note"),
                    value=other,
                )
        self.write_cell(
            ws=ws,
            cell="G50",
            value=self.partner_id.currency_id.with_context(
                lang=self.env.user.lang or "th_TH"
            ).amount_to_text(sum(self.withholding_line_ids.mapped("amount"))),
        )
        self.write_cell(
            ws=ws,
            cell="G63",
            value=self.get_report_locale_date(
                dt=self.document_date,
                py_format="%d %B %Y",
                iso_code="th_TH",
                be_year=True,
            ),
        )
        self.write_cell(
            ws=ws,
            cell=self._get_ws_wht_payment(self.wht_payment),
            value="✓",
        )
        if self.is_stamp_signature:
            (
                self.insert_signature(ws, "H59", self.create_uid.sign_signature)
                if getattr(self.create_uid, "sign_signature", False)
                else self.write_cell(
                    ws=ws,
                    cell="G61",
                    value=html2plaintext(self.create_uid.signature),
                )
            )

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        file_data = base64.b64encode(buffer.read())
        attachment = self.env["ir.attachment"].create(
            {
                "name": "withholding_tax.xlsx",
                "type": "binary",
                "datas": file_data,
                "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",  # noqa: E501
            }
        )
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{attachment.id}?download=true",
            "target": "self",
        }

    def write_cell(
        self,
        ws,
        cell,
        value,
        font=None,
    ):
        c = ws[cell]
        c.value = value
        c.font = font or self._get_style_fount()

    def _get_pnd_sequence_code(self, pnd: str):
        return {
            "pnd3": "account.pnd3",
            "pnd53": "account.pnd53",
        }.get(pnd)

    def _get_pnd_sequence(self, pnd: str):
        sequence_code = self._get_pnd_sequence_code(pnd)
        if not sequence_code:
            return None
        sequence = self.env["ir.sequence"].search(
            [
                ("code", "=", sequence_code),
                ("company_id", "in", [False, self.company_id.id]),
            ]
        )
        if not sequence:
            sequence = self.env["ir.sequence"].create(
                self._prepare_pnd_sequence_vals(sequence_code, pnd)
            )
        return sequence.next_by_id()

    def _get_style_fount(self):
        return Font(
            name="Sarabun",
            size=10,
            color="000000",
        )

    def _get_pnd_type(self, pnd: str) -> str:
        return {
            "pnd3": "H18",
            "pnd53": "M18",
        }.get(pnd)

    def _get_ws_wht_payment(self, wht_payment):
        return {
            "wht": "B57",
            "forever": "B59",
            "once": "B61",
            "other": "B63",
        }.get(wht_payment)

    def _get_tax_section_line(self, tax):
        rows = {
            "1": 23,
            "2": 24,
            "3": 25,
            "4": 26,
            "411": 30,
            "412": 31,
            "413": 32,
            "414": 33,
            "421": 35,
            "422": 36,
            "423": 38,
            "424": 40,
            "425": 41,
            "5": 42,
            "6": 46,
        }
        row = rows.get(tax)
        if not row:
            return None
        data = {
            "date": f"M{row}",
            "base_amount": f"O{row}",
            "amount": f"Q{row}",
        }
        if tax == "425":
            data["note"] = f"G{row}"
        if tax == "6":
            data["note"] = f"E{row}"
        return data

    def insert_signature(self, ws, cell, binary_image):
        if not binary_image:
            return

        image_bytes = base64.b64decode(binary_image)
        image_stream = BytesIO(image_bytes)

        img = XLImage(image_stream)
        img.width = 120
        img.height = 60

        ws.add_image(img, cell)

    def action_open_payment(self):
        self.ensure_one()
        return self.payment_id._get_records_action(name=self.env._("Payment"))

    def action_open_bill(self):
        self.ensure_one()
        return self.move_ids._get_records_action(name=self.env._("Bill"))

    def _prepare_pnd_sequence_vals(self, sequence_code, pnd, **optional):
        return {
            "name": f"{pnd.upper()}",
            "code": sequence_code,
            "padding": 4,
            "use_date_range": True,
            "implementation": "standard",
            "company_id": self.company_id.id,
            "prefix": "%(month)s",
            **optional,
        }

    @api.ondelete(at_uninstall=False)
    def _unlink_check_state(self):
        for rec in self:
            if rec.state == "done":
                raise ValidationError(
                    self.env._(
                        f"You cannot delete a completed withholding tax {rec.name}."
                    )
                )

    @api.model
    def get_report_locale_date(
        self,
        dt: date | datetime,
        py_format: str = "",
        be_year: bool = False,
        iso_code: str = "",
    ) -> str:
        return get_locale_date_format(
            self.env,
            dt,
            py_format,
            be_year,
            iso_code,
        )
